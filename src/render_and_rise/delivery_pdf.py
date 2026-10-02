from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE


def repo_root() -> Path:
    # This file is src/render_and_rise/delivery_pdf.py -> repository root is 2 levels up.
    return Path(__file__).resolve().parents[2]


def _iter_paragraphs(doc: Document):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def _set_run_style(run, *, color_hex: str = "0563C1", underline: bool = True) -> None:
    """
    Word's default hyperlink blue is close to 0563C1. Underline on.
    """
    rPr = run._r.get_or_add_rPr()
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color_hex)
    rPr.append(c)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single" if underline else "none")
    rPr.append(u)


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    """
    Add a real clickable hyperlink into a paragraph.

    python-docx does not expose hyperlink creation, so we create the XML manually.
    """
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    # style: blue + underline
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rPr.append(c)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _clear_paragraph(paragraph) -> None:
    # remove all runs text cleanly (keeps paragraph itself)
    for r in paragraph.runs:
        r.text = ""


def replace_placeholders(doc: Document, title: str, link: str) -> None:
    """
    Template placeholders supported:
      {{TITLE}}
      {{DROPBOX_LINK}}   (preferred)
      {{LINK}}           (legacy)

    For the link placeholder, we insert a real hyperlink with nice display text.
    """
    LINK_TOKENS = ("{{DROPBOX_LINK}}", "{{LINK}}")

    for p in _iter_paragraphs(doc):
        # Title replacement (plain text ok)
        if "{{TITLE}}" in p.text:
            full = p.text.replace("{{TITLE}}", title)
            _clear_paragraph(p)
            if p.runs:
                p.runs[0].text = full
            else:
                p.add_run(full)

        # Link replacement (real hyperlink)
        if any(tok in p.text for tok in LINK_TOKENS):
            # Replace entire paragraph content with: "View your download folder"
            _clear_paragraph(p)

            # Clean UX: not showing the raw URL
            display = "View your download folder"

            # Add the hyperlink run
            _add_hyperlink(p, link, display)


def make_delivery_pdf(
    product: str,
    title: str,
    share_link: str,
    template_docx: Path | None = None,
) -> Path:
    """
    Generates:
      products/<product>/delivery/<Title>_Delivery.pdf

    Requires Word installed (docx2pdf uses Word on Windows).
    """
    root = repo_root()
    product_dir = root / "products" / product
    delivery_dir = product_dir / "delivery"
    delivery_dir.mkdir(parents=True, exist_ok=True)

    # Support a root-level local template, with assets/ as the documented fallback.
    template = template_docx or (root / "delivery_template.docx")
    if not template.exists():
        alt = root / "assets" / "delivery_template.docx"
        if alt.exists():
            template = alt
        else:
            raise FileNotFoundError(f"Missing template DOCX: {template} (and {alt})")

    out_docx = delivery_dir / f"{title}_Delivery.docx"
    out_pdf = delivery_dir / f"{title}_Delivery.pdf"

    doc = Document(str(template))

    # IMPORTANT: share_link should already be dl=0 (view). We do not force dl=1 here.
    replace_placeholders(doc, title=title, link=share_link)

    doc.save(str(out_docx))

    try:
        from docx2pdf import convert
        convert(str(out_docx), str(out_pdf))
    except Exception:
        raise RuntimeError("DOCX->PDF conversion failed. Ensure Microsoft Word is installed.")

    return out_pdf
